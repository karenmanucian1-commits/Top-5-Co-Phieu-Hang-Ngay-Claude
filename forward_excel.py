"""
BÁO CÁO KIỂM ĐỊNH CUỐI NGÀY (Excel) – chấm từng mã V1 / V2 khuyên NÊN MUA bằng giá thật.

  python forward_excel.py             # tạo file Excel trong data/reports/
  python forward_excel.py --upload    # tạo + tải lên Google Drive (secret GDRIVE_WEBHOOK)

Cách tính (giống người mua thật):
  - Lượt 6:00 (dùng giá đóng cửa hôm trước): mua giá MỞ CỬA phiên kế tiếp  -> phiên đó là T0.
  - Lượt 11:35 (dùng giá phiên sáng): mua giá lúc bot báo (giá phiên sáng) -> chính phiên đó là T0.
  - T0, T1 ... T20 = lãi/lỗ theo giá ĐÓNG CỬA của từng phiên so với giá mua, đã trừ phí mua+bán 0.4%.
  - Một mã được khuyên nhiều lần liền: chỉ tính lần đầu; sau 20 phiên khuyên lại mới tính lệnh mới.
  - Kế hoạch của bot (từ T2 – bán được do T+2.5): giá cao nhất chạm mốc chốt lời 1 -> "Chạm chốt lời";
    giá đóng cửa dưới mức cắt lỗ -> "Dính cắt lỗ". Cái nào tới trước tính cái đó;
    cùng phiên xảy ra cả hai thì tính CẮT LỖ (bảo thủ).
"""
import argparse
import base64
import os
from datetime import datetime
from zoneinfo import ZoneInfo

import numpy as np
import pandas as pd
import requests
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import config
import report
import vci_client as vci

D = config.DATA_DIR
LOGS = {"V1": os.path.join(D, "picks_log.csv"), "V2": os.path.join(D, "picks_log_v2.csv")}
HOLD = 20
FEE = config.FEE_ROUNDTRIP
VN_TZ = ZoneInfo("Asia/Ho_Chi_Minh")
TCOLS = [f"T{k}" for k in range(HOLD + 1)]
LUOT = {"CHIEU": "6:00", "SANG": "11:35"}


def load(bot):
    p = LOGS[bot]
    if not os.path.exists(p):
        return pd.DataFrame()
    df = pd.read_csv(p)
    df = df[df["status"] == "MUA"].copy()          # chỉ mã bot khuyên NÊN MUA
    df["run_date"] = pd.to_datetime(df["run_date"])
    df["bot"] = bot
    return df


def entry_info(r, px, days):
    """Trả về (vị trí T0 trong px, giá mua) hoặc None."""
    d = r["run_date"]
    if r["session"] == "SANG":
        same = px.index[px["date"] == d]
        if not len(same):
            return None
        return int(same[0]), float(r["price"])
    after = px.index[px["date"] > d]
    if not len(after):
        return None
    i = int(after[0])
    return i, float(px["open"].iloc[i]) or float(px["close"].iloc[i])


def build_rows(prices, idx):
    days = list(idx["date"])
    pos = {d: i for i, d in enumerate(days)}
    rows = []
    for bot in ("V1", "V2"):
        df = load(bot)
        if df.empty:
            continue
        # sắp theo thời điểm khuyên: cùng ngày thì 11:35 (SANG) trước 6:00 hôm sau (CHIEU cùng run_date)
        df["_o"] = df["session"].map({"SANG": 0, "CHIEU": 1})
        df = df.sort_values(["run_date", "_o", "rank"])
        last, last_i, repeat, kept = {}, {}, {}, []
        for _, r in df.iterrows():
            s, p = r["symbol"], pos.get(r["run_date"])
            if p is None:
                continue
            if s in last and p - last[s] < HOLD:
                repeat[last_i[s]] += 1
                continue
            last[s] = p
            last_i[s] = len(kept)
            repeat[len(kept)] = 0
            kept.append({"key": len(kept), "r": r})
        for k in kept:
            r = k["r"]
            px = prices.get(r["symbol"])
            if px is None or px.empty:
                continue
            px = px.reset_index(drop=True)
            ei = entry_info(r, px, days)
            if ei is None:
                continue
            i0, entry = ei
            seg = px.iloc[i0:i0 + HOLD + 1].reset_index(drop=True)
            n = len(seg)
            row = {
                "Bot": bot,
                "Lượt": LUOT.get(r["session"], r["session"]),
                "Ngày dữ liệu KN": r["run_date"].date(),
                "Ngày mua (T0)": seg["date"].iloc[0].date(),
                "Mã": r["symbol"],
                "Ưu tiên #1": "⭐" if int(r.get("priority", 0) or 0) == 1 else "",
                "Giá mua": round(entry, 2),
                "Mốc chốt lời 1": r.get("t1"),
                "Cắt lỗ": r.get("stop"),
                "Số lần KN lại": repeat[k["key"]],
            }
            for j in range(HOLD + 1):
                row[f"T{j}"] = (seg["close"].iloc[j] / entry - 1 - FEE) * 100 if j < n else np.nan
            row["Hiện tại %"] = (seg["close"].iloc[-1] / entry - 1 - FEE) * 100
            row["Cao nhất %"] = (seg["high"].max() / entry - 1) * 100
            row["Thấp nhất %"] = (seg["low"].min() / entry - 1) * 100
            # VN-Index cùng kỳ (từ giá mở cửa T0 đến phiên hiện tại của lệnh)
            t0 = seg["date"].iloc[0]
            ia = idx[idx["date"] >= t0].reset_index(drop=True)
            if len(ia):
                ie = float(ia["open"].iloc[0]) or float(ia["close"].iloc[0])
                m = min(len(ia), n)
                row["VNI cùng kỳ %"] = (ia["close"].iloc[m - 1] / ie - 1) * 100
                row["Hơn VNI %"] = row["Hiện tại %"] - row["VNI cùng kỳ %"]
                for j in range(min(m, HOLD + 1)):
                    row[f"_vni{j}"] = (ia["close"].iloc[j] / ie - 1) * 100
            plan = "Chưa chạm"
            t1, stop = r.get("t1"), r.get("stop")
            for j in range(2, n):
                # Bảo thủ: cùng phiên vừa chạm T1 vừa đóng cửa dưới cắt lỗ -> tính CẮT LỖ trước
                if pd.notna(stop) and stop and seg["close"].iloc[j] < stop:
                    plan = f"Dính cắt lỗ (T{j})"
                    break
                if pd.notna(t1) and t1 and seg["high"].iloc[j] >= t1:
                    plan = f"Chạm chốt lời (T{j})"
                    break
            row["Kế hoạch bot"] = plan
            row["Trạng thái"] = "Đã đủ 20 phiên" if n > HOLD else f"Đang theo dõi ({n - 1}/20 phiên)"
            rows.append(row)
    return pd.DataFrame(rows)


# ---------------------------------------------------------------- Excel
HEAD = PatternFill("solid", fgColor="1F4E78")
SUB = PatternFill("solid", fgColor="DDEBF7")
GREEN = PatternFill("solid", fgColor="C6EFCE")
RED = PatternFill("solid", fgColor="FFC7CE")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _header(ws, row, cols):
    for c, name in enumerate(cols, 1):
        cell = ws.cell(row=row, column=c, value=name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEAD
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BOX


def _color_pct(ws, rng):
    ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["0"], fill=GREEN,
                                                  font=Font(color="006100")))
    ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0"], fill=RED,
                                                  font=Font(color="9C0006")))


def sheet_detail(wb, res, title):
    ws = wb.create_sheet(title)
    cols = ["Bot", "Lượt", "Ngày mua (T0)", "Mã", "Ưu tiên #1", "Giá mua", "Mốc chốt lời 1", "Cắt lỗ",
            *TCOLS, "Hiện tại %", "Cao nhất %", "Thấp nhất %", "VNI cùng kỳ %", "Hơn VNI %",
            "Kế hoạch bot", "Trạng thái", "Số lần KN lại"]
    _header(ws, 1, cols)
    if not res.empty:
        res = res.sort_values(["Ngày mua (T0)", "Bot", "Mã"], ascending=[False, True, True])
    for i, (_, r) in enumerate(res.iterrows(), 2):
        for c, name in enumerate(cols, 1):
            v = r.get(name)
            if isinstance(v, float) and np.isnan(v):
                v = None
            cell = ws.cell(row=i, column=c, value=v)
            cell.border = BOX
            if name in TCOLS or name.endswith("%"):
                cell.number_format = '+0.0"%";-0.0"%";0.0"%"'
            elif name in ("Giá mua", "Mốc chốt lời 1", "Cắt lỗ"):
                cell.number_format = "#,##0.00"
            elif name == "Ngày mua (T0)":
                cell.number_format = "dd/mm/yyyy"
            if name == "Kế hoạch bot" and isinstance(v, str):
                cell.fill = GREEN if v.startswith("Chạm") else RED if v.startswith("Dính") else PatternFill()
    last = ws.max_row
    if last >= 2:
        a = get_column_letter(cols.index("T0") + 1)
        b = get_column_letter(cols.index("Hơn VNI %") + 1)
        _color_pct(ws, f"{a}2:{b}{last}")
    widths = {"Bot": 5, "Lượt": 6, "Ngày mua (T0)": 11, "Mã": 6, "Ưu tiên #1": 6, "Giá mua": 9,
              "Mốc chốt lời 1": 9, "Cắt lỗ": 8, "Kế hoạch bot": 20, "Trạng thái": 22, "Số lần KN lại": 8}
    for c, name in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(c)].width = widths.get(name, 7.5)
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "F2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(last, 1)}"
    return ws


def sheet_summary(wb, res, now):
    ws = wb.active
    ws.title = "Tổng hợp"
    ws["A1"] = f"BÁO CÁO KIỂM ĐỊNH CLAUDE BOT V1 & V2 – {now:%H:%M %d/%m/%Y}"
    ws["A1"].font = Font(bold=True, size=14, color="1F4E78")
    ws["A2"] = ("Chỉ chấm mã bot khuyên NÊN MUA. Lượt 6:00 mua giá mở cửa phiên T0; lượt 11:35 mua giá phiên sáng. "
                "T0…T20 = lãi/lỗ theo giá đóng cửa từng phiên, đã trừ phí 0.4%. Mỗi mã chỉ tính lần khuyên đầu tiên.")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:L2")
    ws.row_dimensions[2].height = 45

    r0 = 4
    if res.empty:
        ws.cell(row=r0, column=1, value="Chưa có lệnh nào có giá sau ngày khuyến nghị.")
        return ws

    # Bảng 1: tổng quan mỗi bot
    cols = ["Bot", "Tổng lệnh", "Đang theo dõi", "Đủ 20 phiên", "Lãi/lỗ hiện tại TB %",
            "% lệnh đang lãi", "VNI cùng kỳ TB %", "Hơn VNI TB %", "Chạm chốt lời", "Dính cắt lỗ", "Chưa chạm"]
    ws.cell(row=r0, column=1, value="1. TỔNG QUAN").font = Font(bold=True)
    _header(ws, r0 + 1, cols)
    rr = r0 + 2
    for bot in ("V1", "V2"):
        g = res[res["Bot"] == bot]
        if g.empty:
            continue
        plan = g["Kế hoạch bot"].str.split(" ").str[0]
        vals = [bot, len(g), int((~g["Trạng thái"].str.startswith("Đã")).sum()),
                int(g["Trạng thái"].str.startswith("Đã").sum()), g["Hiện tại %"].mean(),
                (g["Hiện tại %"] > 0).mean() * 100, g["VNI cùng kỳ %"].mean(), g["Hơn VNI %"].mean(),
                int((plan == "Chạm").sum()), int((plan == "Dính").sum()), int((plan == "Chưa").sum())]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=rr, column=c, value=v)
            cell.border = BOX
            if isinstance(v, float):
                cell.number_format = '+0.00"%";-0.00"%";0.00"%"' if c != 6 else '0"%"'
        rr += 1
    _color_pct(ws, f"E{r0 + 2}:E{rr - 1}")
    _color_pct(ws, f"H{r0 + 2}:H{rr - 1}")

    # Bảng 2: lãi/lỗ trung bình theo từng phiên T0..T20
    r1 = rr + 2
    ws.cell(row=r1, column=1, value="2. LÃI/LỖ TRUNG BÌNH THEO TỪNG PHIÊN SAU KHI MUA").font = Font(bold=True)
    cols2 = ["Phiên", "V1 số lệnh", "V1 lãi TB %", "V1 % lệnh lãi", "V1 hơn VNI %",
             "V2 số lệnh", "V2 lãi TB %", "V2 % lệnh lãi", "V2 hơn VNI %"]
    _header(ws, r1 + 1, cols2)
    rr = r1 + 2
    for j, t in enumerate(TCOLS):
        vals = [t]
        for bot in ("V1", "V2"):
            g = res[(res["Bot"] == bot) & res[t].notna()]
            if len(g):
                vni = g.get(f"_vni{j}")
                vals += [len(g), g[t].mean(), (g[t] > 0).mean() * 100,
                         (g[t] - vni).mean() if vni is not None else None]
            else:
                vals += [0, None, None, None]
        for c, v in enumerate(vals, 1):
            if isinstance(v, float) and np.isnan(v):
                v = None
            cell = ws.cell(row=rr, column=c, value=v)
            cell.border = BOX
            if c in (3, 5, 7, 9):
                cell.number_format = '+0.00"%";-0.00"%";0.00"%"'
            elif c in (4, 8):
                cell.number_format = '0"%"'
        rr += 1
    for col in "CEGI":
        _color_pct(ws, f"{col}{r1 + 2}:{col}{rr - 1}")
    ws.cell(row=rr + 1, column=1,
            value="Ghi chú: chưa đủ ~30 lệnh đủ 20 phiên thì kết quả còn ít mẫu, chưa nên kết luận bot nào tốt hơn."
            ).font = Font(italic=True, color="7F7F7F")
    for c in range(1, 12):
        ws.column_dimensions[get_column_letter(c)].width = 14
    ws.column_dimensions["A"].width = 10
    return ws


def make_excel(res, now, path):
    wb = Workbook()
    sheet_summary(wb, res, now)
    sheet_detail(wb, res, "Chi tiết từng mã")
    for bot in ("V1", "V2"):
        sheet_detail(wb, res[res["Bot"] == bot] if not res.empty else res, f"Chi tiết {bot}")
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)
    return path


def upload(path):
    url = os.getenv("GDRIVE_WEBHOOK")
    if not url:
        print("Chưa có GDRIVE_WEBHOOK -> bỏ qua tải lên Google Drive")
        return False
    with open(path, "rb") as f:
        data = base64.b64encode(f.read()).decode()
    r = requests.post(url, json={"name": os.path.basename(path), "data": data}, timeout=120)
    print("Google Drive:", r.status_code, r.text[:200])
    if not r.ok or not r.text.startswith("OK"):
        raise RuntimeError("Tải lên Google Drive thất bại")
    return True


def run(do_upload=False, client=vci, now=None):
    now = now or datetime.now(VN_TZ)
    syms = set()
    for bot in LOGS:
        df = load(bot)
        if len(df):
            syms |= set(df["symbol"])
    idx = client.get_ohlc([config.INDEX_SYMBOL], bars=150)[config.INDEX_SYMBOL]
    idx["date"] = pd.to_datetime(idx["date"])
    prices = client.get_ohlc(sorted(syms), bars=150) if syms else {}
    for s in prices:
        prices[s]["date"] = pd.to_datetime(prices[s]["date"])
    res = build_rows(prices, idx)
    name = "Kiem dinh CLAUDE BOT.xlsx"          # 1 file duy nhất, cập nhật đè mỗi ngày
    path = make_excel(res, now, os.path.join(D, "reports", name))
    res.drop(columns=[c for c in res.columns if c.startswith("_")], errors="ignore") \
       .to_csv(os.path.join(D, "forward_test.csv"), index=False, encoding="utf-8-sig")
    print(f"Đã tạo {path} – {len(res)} lệnh")
    if do_upload:
        try:
            ok = upload(path)
        except Exception as e:
            print("Lỗi tải lên Drive:", e)
            ok = False
        if ok:
            msg = ("📊 FORWARD TEST _ CLAUDE\n"
                   f"✅ Đã cập nhật Báo cáo kiểm định ngày {now:%d/%m/%Y}.\n"
                   "📁 Báo cáo chi tiết đã được cập nhật trên Google Drive.")
        else:
            msg = ("📊 FORWARD TEST _ CLAUDE\n"
                   f"⚠️ Chưa tải được Báo cáo kiểm định ngày {now:%d/%m/%Y} lên Google Drive.")
        report.send_telegram([msg])
        if not ok:
            raise SystemExit(1)
    return res


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--upload", action="store_true")
    run(do_upload=ap.parse_args().upload)
